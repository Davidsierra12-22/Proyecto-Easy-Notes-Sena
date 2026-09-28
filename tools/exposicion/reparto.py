# -*- coding: utf-8 -*-
"""Resumen del reparto real: lee las notas del PPTX y calcula tiempos y porcentajes."""
import re
import sys
from pptx import Presentation

PPM = 135
patron = re.compile(r'^([A-ZÁÉÍÓÚÑ]+) \((?:(\d+)m )?(\d+)\s*s\)\.\s*')

ruta = sys.argv[1] if len(sys.argv) > 1 else 'docs/PLAN DE NEGOCIO - EXPOSICION.pptx'
prs = Presentation(ruta)

agg = {}
orden = []
malas = []
for i, s in enumerate(prs.slides, 1):
    t = s.notes_slide.notes_text_frame.text
    m = patron.match(t)
    if not m:
        malas.append(i)
        continue
    quien, mm, ss = m.group(1), m.group(2), m.group(3)
    seg = (int(mm) * 60 if mm else 0) + int(ss)
    d = agg.setdefault(quien, {'slides': [], 'pal': 0, 'seg': 0})
    d['slides'].append(i)
    d['pal'] += len(t.split())
    d['seg'] += seg
    if quien not in orden:
        orden.append(quien)

if malas:
    print('Sin etiqueta de tiempo en:', malas)

tot_seg = sum(d['seg'] for d in agg.values())
tot_pal = sum(d['pal'] for d in agg.values())
print()
print(f"{'Orador':10} {'Diapositivas':14} {'Palabras':>8} {'Tiempo':>8} {'% tiempo':>9}")
print('-' * 53)
for q in orden:
    d = agg[q]
    print(f"{q:10} {','.join(map(str, d['slides'])):14} {d['pal']:>8} "
          f"{d['seg'] // 60}m{d['seg'] % 60:02d}s {d['seg'] / tot_seg * 100:>8.1f}%")
print('-' * 53)
print(f"{'TOTAL':10} {'1-17':14} {tot_pal:>8} "
      f"{tot_seg // 60}m{tot_seg % 60:02d}s {100.0:>8.1f}%")
print(f"\n({tot_pal} palabras a {PPM} palabras/min)")
