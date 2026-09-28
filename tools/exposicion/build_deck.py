# -*- coding: utf-8 -*-
"""Compila el deck de exposicion de EasyNotes v2 (17 diapositivas)."""

import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.util import Inches
import deck_slides as D

OUT = '/home/lenovot440/Escritorio/easynots-modelos/docs/PLAN DE NEGOCIO - EXPOSICION.pptx'

ORDEN = [
    D.s01_portada, D.s02_problema, D.s03_solucion, D.s04_objetivos,
    D.s05_mercado, D.s06_competencia, D.s07_producto, D.s08_precios,
    D.s09_economia_unitaria, D.s10_canales, D.s11_proyeccion,
    D.s12_gantt, D.s13_operacion, D.s14_numeros, D.s15_dofa,
    D.s16_impacto, D.s17_cierre,
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
for fn in ORDEN:
    fn(prs)

cp = prs.core_properties
cp.title = 'EasyNotes v2 - Plan de Negocio Básico 2026'
cp.author = 'Martín Zapata, Sneyder Gómez, Yorman Gómez, Santiago Sierra, Andrés Avila'
cp.subject = 'Exposición del plan de negocio (15-20 minutos)'
cp.comments = ('Las notas de cada diapositiva contienen el guion del orador, el ponente '
               'y el tiempo asignado.')

prs.save(OUT)

print('Diapositivas: %d' % len(ORDEN))
print()
print('%-4s %-9s %-8s %6s' % ('#', 'ponente', 'seccion', 'seg'))
print('-' * 34)
total = 0
for i, s in enumerate(prs.slides, 1):
    txt = s.notes_slide.notes_text_frame.text
    palabras = len(txt.split())
    total += palabras
    ponente = txt.split(' ')[0].split('(')[0].strip().strip('—').strip()
    seccion = ''
    for sh in s.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t and t[0].isdigit() and len(t) < 46:
                seccion = t
                break
    print('%-4d %-9s %-8s %6d' % (i, ponente[:9], seccion[:8],
                                 round(palabras / 2.5)))
print('-' * 34)
print('Guion: %d palabras | ~%d min a 150 pal/min | ~%d min a 135 pal/min'
      % (total, round(total / 150), round(total / 135)))
print('\nGuardado:', OUT)
