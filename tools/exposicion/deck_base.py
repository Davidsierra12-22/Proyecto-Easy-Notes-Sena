# -*- coding: utf-8 -*-
"""Primitivas de maquetacion del deck de exposicion de EasyNotes v2."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

AZUL = RGBColor(0x1F, 0x4E, 0x79)
AZUL_CLARO = RGBColor(0x2E, 0x74, 0xB5)
GRIS = RGBColor(0x40, 0x40, 0x40)
GRIS_CLARO = RGBColor(0x7F, 0x7F, 0x7F)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
VERDE = RGBColor(0x2E, 0x7D, 0x32)
ROJO = RGBColor(0xB0, 0x30, 0x30)
FONDO = RGBColor(0xF7, 0xF9, 0xFB)

FUENTE = 'Arial'
W, H = Inches(13.333), Inches(7.5)
EQUIPO = ['Martín Zapata', 'Sneyder Gómez', 'Yorman Gómez',
          'Santiago Sierra', 'Andrés Avila']


def nueva_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BLANCO
    return s


def caja(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    return tf


def parrafo(tf, text, size=18, bold=False, color=GRIS, align=PP_ALIGN.LEFT,
            space_after=6, space_before=0, first=False, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    r = p.add_run()
    r.text = text
    r.font.name = FUENTE
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return p


def vinetas(slide, x, y, w, h, items, size=18, color=GRIS, gap=10):
    tf = caja(slide, x, y, w, h)
    for i, it in enumerate(items):
        parrafo(tf, '•  ' + it, size=size, color=color, space_after=gap,
                first=(i == 0))
    return tf


def titulo(slide, seccion, texto, sub=None):
    """Barra superior con el numero de seccion de la plantilla SENA."""
    if seccion:
        barra = slide.shapes.add_shape(1, Inches(0), Inches(0), W, Inches(0.42))
        barra.fill.solid()
        barra.fill.fore_color.rgb = AZUL
        barra.line.fill.background()
        barra.shadow.inherit = False
        tfb = barra.text_frame
        tfb.margin_left = Inches(0.45)
        tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
        parrafo(tfb, seccion, size=12, bold=True, color=BLANCO, first=True,
                space_after=0)
    tf = caja(slide, Inches(0.45), Inches(0.6), Inches(12.45), Inches(1.0))
    parrafo(tf, texto, size=28, bold=True, color=AZUL, first=True, space_after=0)
    if sub:
        parrafo(tf, sub, size=14, color=GRIS_CLARO, space_before=4)


def pie(slide, texto, color=GRIS_CLARO):
    tf = caja(slide, Inches(0.45), Inches(6.95), Inches(12.45), Inches(0.35))
    parrafo(tf, texto, size=10, color=color, first=True, space_after=0)


# Ritmo de lectura assumed para exposición (palabras por minuto).
PPM = 135


def duracion(texto):
    """Devuelve la duración real de un guion, calculada sobre su texto.

    Se calcula siempre, nunca a mano: así la etiqueta de tiempo no puede
    quedar desalineada del guion que realmente se va a leer.
    """
    seg = round(len(texto.split()) / PPM * 60)
    return f'{seg // 60}m {seg % 60:02d}s' if seg >= 60 else f'{seg} s'


def notas(slide, quien, texto):
    """Escribe el guion del orador, prefijando el nombre y el tiempo real.

    El texto que se pasa es el cuerpo del guion, sin la etiqueta: la etiqueta
    la arma esta función, de modo que nombre y duración siempre corresponden.
    """
    slide.notes_slide.notes_text_frame.text = f'{quien} ({duracion(texto)}). {texto.strip()}'


def tabla(slide, x, y, w, h, datos, anchos=None, size=12, header=True,
          alto_fila=0.32, aligns=None):
    """datos: lista de filas; la primera es encabezado si header=True."""
    rows, cols = len(datos), len(datos[0])
    gt = slide.shapes.add_table(rows, cols, x, y, w, h).table
    if anchos:
        total = sum(anchos)
        for i, a in enumerate(anchos):
            gt.columns[i].width = Emu(int(w * a / total))
    aligns = aligns or [None] * cols
    for ri, fila in enumerate(datos):
        gt.rows[ri].height = Inches(alto_fila if ri else alto_fila + 0.06)
        for ci, val in enumerate(fila):
            c = gt.cell(ri, ci)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            es_head = ri == 0 and header
            c.fill.fore_color.rgb = AZUL if es_head else (BLANCO if ri % 2 else FONDO)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (aligns[ci] or PP_ALIGN.LEFT)
            r = p.add_run()
            r.text = str(val)
            r.font.name = FUENTE
            r.font.size = Pt(size)
            r.font.bold = es_head
            r.font.color.rgb = BLANCO if es_head else GRIS
    return gt


def cifra(slide, x, y, w, valor, etiqueta, color=AZUL, h=1.15, size_val=22):
    """Tarjeta de cifra destacada."""
    box = slide.shapes.add_shape(1, x, y, w, Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = FONDO
    box.line.color.rgb = color
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.06)
    parrafo(tf, valor, size=size_val, bold=True, color=color,
            align=PP_ALIGN.CENTER, first=True, space_after=2)
    parrafo(tf, etiqueta, size=10.5, color=GRIS, align=PP_ALIGN.CENTER,
            space_after=0)
    return box


def panel(slide, x, y, w, h, titulo_txt, lineas, color=AZUL, size=14,
          size_titulo=12):
    """Recuadro con titulo y lista de puntos."""
    box = slide.shapes.add_shape(1, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = FONDO
    box.line.color.rgb = color
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12)
    parrafo(tf, titulo_txt, size=size_titulo, bold=True, color=color,
            first=True, space_after=6)
    for ln in lineas:
        parrafo(tf, ln, size=size, color=GRIS, space_after=6)
    return box
